import io
import re
from typing import Dict, Any, List, Optional
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable, KeepTogether
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT, TA_JUSTIFY
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT


class ResumeExportService:
    @classmethod
    def generate_pdf(
        cls,
        resume_data: Dict[str, Any],
        template_name: str = "modern_clean"
    ) -> bytes:
        """Generate ATS-friendly PDF using ReportLab."""
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=letter,
            leftMargin=36,
            rightMargin=36,
            topMargin=36,
            bottomMargin=36
        )

        styles = getSampleStyleSheet()
        
        # Color palettes
        is_modern = (template_name == "modern_clean")
        primary_color = colors.HexColor("#1e3a8a") if is_modern else colors.HexColor("#111827")
        accent_color = colors.HexColor("#2563eb") if is_modern else colors.HexColor("#374151")
        text_color = colors.HexColor("#1f2937")
        muted_color = colors.HexColor("#4b5563")
        line_color = colors.HexColor("#cbd5e1") if is_modern else colors.HexColor("#9ca3af")

        font_name = "Helvetica" if is_modern else "Times-Roman"
        font_bold = "Helvetica-Bold" if is_modern else "Times-Bold"
        font_italic = "Helvetica-Oblique" if is_modern else "Times-Italic"

        # Custom Paragraph Styles
        name_style = ParagraphStyle(
            "ResumeName",
            parent=styles["Normal"],
            fontName=font_bold,
            fontSize=18 if is_modern else 20,
            leading=22,
            alignment=TA_LEFT if is_modern else TA_CENTER,
            textColor=primary_color
        )

        headline_style = ParagraphStyle(
            "ResumeHeadline",
            parent=styles["Normal"],
            fontName=font_bold if is_modern else font_italic,
            fontSize=11,
            leading=14,
            alignment=TA_LEFT if is_modern else TA_CENTER,
            textColor=accent_color
        )

        contact_style = ParagraphStyle(
            "ResumeContact",
            parent=styles["Normal"],
            fontName=font_name,
            fontSize=9,
            leading=12,
            alignment=TA_LEFT if is_modern else TA_CENTER,
            textColor=muted_color
        )

        section_heading_style = ParagraphStyle(
            "ResumeSectionHeading",
            parent=styles["Normal"],
            fontName=font_bold,
            fontSize=11,
            leading=14,
            textColor=primary_color,
            spaceBefore=8,
            spaceAfter=2,
            textTransform="uppercase" if not is_modern else None
        )

        body_style = ParagraphStyle(
            "ResumeBody",
            parent=styles["Normal"],
            fontName=font_name,
            fontSize=9.5,
            leading=13,
            textColor=text_color
        )

        bullet_style = ParagraphStyle(
            "ResumeBullet",
            parent=styles["Normal"],
            fontName=font_name,
            fontSize=9,
            leading=12.5,
            textColor=text_color,
            leftIndent=14,
            firstLineIndent=-10
        )

        role_company_style = ParagraphStyle(
            "ResumeRoleCompany",
            parent=styles["Normal"],
            fontName=font_bold,
            fontSize=10,
            leading=13,
            textColor=text_color
        )

        date_loc_style = ParagraphStyle(
            "ResumeDateLoc",
            parent=styles["Normal"],
            fontName=font_italic if not is_modern else font_name,
            fontSize=9,
            leading=13,
            alignment=TA_RIGHT,
            textColor=muted_color
        )

        story = []
        profile = resume_data.get("profile", {})

        # 1. Header (Name & Headline)
        story.append(Paragraph(profile.get("full_name", "Candidate Name"), name_style))
        if profile.get("headline"):
            story.append(Paragraph(profile.get("headline", ""), headline_style))

        # Contact details bar
        contacts = []
        if profile.get("email"):
            contacts.append(profile["email"])
        if profile.get("phone"):
            contacts.append(profile["phone"])
        if profile.get("location"):
            contacts.append(profile["location"])
        if profile.get("linkedin_url"):
            contacts.append(profile["linkedin_url"].replace("https://", ""))
        if profile.get("github_url"):
            contacts.append(profile["github_url"].replace("https://", ""))

        story.append(Spacer(1, 2))
        story.append(Paragraph(" &nbsp;|&nbsp; ".join(contacts), contact_style))
        story.append(Spacer(1, 4))
        story.append(HRFlowable(width="100%", thickness=1, color=line_color, spaceBefore=2, spaceAfter=6))

        # 2. Professional Summary
        summary = resume_data.get("summary") or profile.get("summary")
        if summary:
            story.append(Paragraph("PROFESSIONAL SUMMARY", section_heading_style))
            story.append(HRFlowable(width="100%", thickness=0.5, color=line_color, spaceBefore=1, spaceAfter=4))
            story.append(Paragraph(summary, body_style))
            story.append(Spacer(1, 4))

        # 3. Technical Skills
        skills = resume_data.get("skills", [])
        if skills:
            story.append(Paragraph("TECHNICAL & CORE SKILLS", section_heading_style))
            story.append(HRFlowable(width="100%", thickness=0.5, color=line_color, spaceBefore=1, spaceAfter=4))

            # Group skills by category
            tech_s = [s["name"] for s in skills if s.get("category") == "technical"]
            tool_s = [s["name"] for s in skills if s.get("category") == "tool"]
            soft_s = [s["name"] for s in skills if s.get("category") == "soft"]
            other_s = [s["name"] for s in skills if s.get("category") not in ("technical", "tool", "soft")]

            skill_rows = []
            if tech_s:
                skill_rows.append(Paragraph(f"<b>Languages & Frameworks:</b> {', '.join(tech_s)}", body_style))
            if tool_s:
                skill_rows.append(Paragraph(f"<b>Tools & Infrastructure:</b> {', '.join(tool_s)}", body_style))
            if soft_s:
                skill_rows.append(Paragraph(f"<b>Architecture & Leadership:</b> {', '.join(soft_s)}", body_style))
            if other_s:
                skill_rows.append(Paragraph(f"<b>Additional Competencies:</b> {', '.join(other_s)}", body_style))

            for sr in skill_rows:
                story.append(sr)
            story.append(Spacer(1, 4))

        # 4. Work Experience
        experience = resume_data.get("experience", [])
        if experience:
            story.append(Paragraph("PROFESSIONAL EXPERIENCE", section_heading_style))
            story.append(HRFlowable(width="100%", thickness=0.5, color=line_color, spaceBefore=1, spaceAfter=4))

            for exp in experience:
                title_text = f"<b>{exp.get('role', '')}</b> &mdash; {exp.get('company', '')}"
                date_text = f"{exp.get('start_date', '')} &ndash; {exp.get('end_date', 'Present')}"
                if exp.get("location"):
                    date_text += f" | {exp.get('location')}"

                # Table for 2-column header (Role/Company left, Date/Loc right)
                header_table = Table(
                    [[Paragraph(title_text, role_company_style), Paragraph(date_text, date_loc_style)]],
                    colWidths=[360, 180]
                )
                header_table.setStyle(TableStyle([
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 0),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 0),
                    ("TOPPADDING", (0, 0), (-1, -1), 1),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 1),
                ]))
                story.append(header_table)

                # Bullets
                bullets = exp.get("bullet_points", [])
                for b in bullets:
                    story.append(Paragraph(f"&bull;&nbsp;&nbsp;{b}", bullet_style))
                story.append(Spacer(1, 4))

        # 5. Key Projects
        projects = resume_data.get("projects", [])
        if projects:
            story.append(Paragraph("NOTABLE PROJECTS", section_heading_style))
            story.append(HRFlowable(width="100%", thickness=0.5, color=line_color, spaceBefore=1, spaceAfter=4))

            for proj in projects:
                tech_line = f" ({', '.join(proj.get('tech_stack', []))})" if proj.get("tech_stack") else ""
                title_text = f"<b>{proj.get('title', '')}</b>{tech_line}"
                date_text = f"{proj.get('start_date', '')} &ndash; {proj.get('end_date', 'Present')}" if proj.get("start_date") else ""

                header_table = Table(
                    [[Paragraph(title_text, role_company_style), Paragraph(date_text, date_loc_style)]],
                    colWidths=[370, 170]
                )
                header_table.setStyle(TableStyle([
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 0),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 0),
                    ("TOPPADDING", (0, 0), (-1, -1), 1),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 1),
                ]))
                story.append(header_table)

                for b in proj.get("bullet_points", []):
                    story.append(Paragraph(f"&bull;&nbsp;&nbsp;{b}", bullet_style))
                story.append(Spacer(1, 4))

        # 6. Education
        education = resume_data.get("education", [])
        if education:
            story.append(Paragraph("EDUCATION", section_heading_style))
            story.append(HRFlowable(width="100%", thickness=0.5, color=line_color, spaceBefore=1, spaceAfter=4))

            for edu in education:
                title_text = f"<b>{edu.get('institution', '')}</b> &mdash; {edu.get('degree', '')} in {edu.get('field_of_study', '')}"
                date_text = f"{edu.get('start_date', '')} &ndash; {edu.get('end_date', 'Present')}"
                if edu.get("gpa"):
                    title_text += f" (GPA: {edu.get('gpa')})"

                header_table = Table(
                    [[Paragraph(title_text, role_company_style), Paragraph(date_text, date_loc_style)]],
                    colWidths=[380, 160]
                )
                header_table.setStyle(TableStyle([
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 0),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 0),
                    ("TOPPADDING", (0, 0), (-1, -1), 1),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 1),
                ]))
                story.append(header_table)

                for h in edu.get("highlights", []):
                    story.append(Paragraph(f"&bull;&nbsp;&nbsp;{h}", bullet_style))
                story.append(Spacer(1, 3))

        # 7. Certifications & Achievements
        certs = resume_data.get("certifications", [])
        achievements = resume_data.get("achievements", [])
        if certs or achievements:
            story.append(Paragraph("CERTIFICATIONS & HONORS", section_heading_style))
            story.append(HRFlowable(width="100%", thickness=0.5, color=line_color, spaceBefore=1, spaceAfter=4))

            for c in certs:
                c_text = f"<b>{c.get('name')}</b> &mdash; {c.get('issuer')} ({c.get('issue_date')})"
                story.append(Paragraph(f"&bull;&nbsp;&nbsp;{c_text}", bullet_style))

            for a in achievements:
                a_text = f"<b>{a.get('title')}</b>: {a.get('description')}"
                story.append(Paragraph(f"&bull;&nbsp;&nbsp;{a_text}", bullet_style))

        # Build document
        doc.build(story)
        pdf_bytes = buffer.getvalue()
        buffer.close()
        return pdf_bytes

    @classmethod
    def generate_docx(
        cls,
        resume_data: Dict[str, Any],
        template_name: str = "modern_clean"
    ) -> bytes:
        """Generate ATS-friendly DOCX document."""
        doc = Document()

        # Set 0.5 inch margins
        for section in doc.sections:
            section.top_margin = Inches(0.5)
            section.bottom_margin = Inches(0.5)
            section.left_margin = Inches(0.5)
            section.right_margin = Inches(0.5)

        is_modern = (template_name == "modern_clean")
        font_family = "Arial" if is_modern else "Times New Roman"

        profile = resume_data.get("profile", {})

        # Name
        p_name = doc.add_paragraph()
        p_name.alignment = WD_ALIGN_PARAGRAPH.LEFT if is_modern else WD_ALIGN_PARAGRAPH.CENTER
        run_name = p_name.add_run(profile.get("full_name", "Candidate Name"))
        run_name.bold = True
        run_name.font.size = Pt(20 if not is_modern else 18)
        run_name.font.name = font_family
        if is_modern:
            run_name.font.color.rgb = RGBColor(30, 58, 138)

        # Headline
        if profile.get("headline"):
            p_head = doc.add_paragraph()
            p_head.alignment = WD_ALIGN_PARAGRAPH.LEFT if is_modern else WD_ALIGN_PARAGRAPH.CENTER
            run_head = p_head.add_run(profile.get("headline"))
            run_head.font.size = Pt(11)
            run_head.font.name = font_family
            run_head.italic = not is_modern

        # Contacts
        contacts = []
        if profile.get("email"):
            contacts.append(profile["email"])
        if profile.get("phone"):
            contacts.append(profile["phone"])
        if profile.get("location"):
            contacts.append(profile["location"])
        if profile.get("linkedin_url"):
            contacts.append(profile["linkedin_url"].replace("https://", ""))
        if profile.get("github_url"):
            contacts.append(profile["github_url"].replace("https://", ""))

        p_contact = doc.add_paragraph()
        p_contact.alignment = WD_ALIGN_PARAGRAPH.LEFT if is_modern else WD_ALIGN_PARAGRAPH.CENTER
        run_c = p_contact.add_run(" | ".join(contacts))
        run_c.font.size = Pt(9.5)
        run_c.font.name = font_family
        run_c.font.color.rgb = RGBColor(100, 116, 139)

        def add_heading(title: str):
            h = doc.add_paragraph()
            h.paragraph_format.space_before = Pt(8)
            h.paragraph_format.space_after = Pt(2)
            r = h.add_run(title.upper())
            r.bold = True
            r.font.size = Pt(11)
            r.font.name = font_family
            if is_modern:
                r.font.color.rgb = RGBColor(30, 58, 138)

        # Summary
        summary = resume_data.get("summary") or profile.get("summary")
        if summary:
            add_heading("Professional Summary")
            p_s = doc.add_paragraph(summary)
            p_s.paragraph_format.space_after = Pt(4)
            p_s.style.font.name = font_family
            p_s.style.font.size = Pt(9.5)

        # Skills
        skills = resume_data.get("skills", [])
        if skills:
            add_heading("Technical Skills")
            skill_names = [s["name"] for s in skills]
            p_sk = doc.add_paragraph(f"Skills: {', '.join(skill_names)}")
            p_sk.paragraph_format.space_after = Pt(4)

        # Experience
        experience = resume_data.get("experience", [])
        if experience:
            add_heading("Professional Experience")
            for exp in experience:
                p_exp = doc.add_paragraph()
                r_title = p_exp.add_run(f"{exp.get('role')} — {exp.get('company')}")
                r_title.bold = True
                p_exp.add_run(f" | {exp.get('start_date')} – {exp.get('end_date', 'Present')}")
                p_exp.paragraph_format.space_after = Pt(1)

                for b in exp.get("bullet_points", []):
                    p_b = doc.add_paragraph(b, style="List Bullet")
                    p_b.paragraph_format.space_after = Pt(1)

        # Projects
        projects = resume_data.get("projects", [])
        if projects:
            add_heading("Key Projects")
            for proj in projects:
                p_pr = doc.add_paragraph()
                r_title = p_pr.add_run(proj.get("title"))
                r_title.bold = True
                if proj.get("tech_stack"):
                    p_pr.add_run(f" ({', '.join(proj.get('tech_stack'))})")
                p_pr.paragraph_format.space_after = Pt(1)

                for b in proj.get("bullet_points", []):
                    p_b = doc.add_paragraph(b, style="List Bullet")
                    p_b.paragraph_format.space_after = Pt(1)

        # Education
        education = resume_data.get("education", [])
        if education:
            add_heading("Education")
            for edu in education:
                p_ed = doc.add_paragraph()
                r_inst = p_ed.add_run(edu.get("institution"))
                r_inst.bold = True
                p_ed.add_run(f" — {edu.get('degree')} in {edu.get('field_of_study')} ({edu.get('start_date')} – {edu.get('end_date', 'Present')})")
                p_ed.paragraph_format.space_after = Pt(2)

        buffer = io.BytesIO()
        doc.save(buffer)
        docx_bytes = buffer.getvalue()
        buffer.close()
        return docx_bytes

    @classmethod
    def generate_html_preview(
        cls,
        resume_data: Dict[str, Any],
        template_name: str = "modern_clean"
    ) -> str:
        """Render live responsive HTML representation for the frontend preview."""
        is_modern = (template_name == "modern_clean")
        font_family = "Inter, system-ui, sans-serif" if is_modern else "'Merriweather', Georgia, serif"
        primary_color = "#1e3a8a" if is_modern else "#111827"
        accent_color = "#2563eb" if is_modern else "#4b5563"
        border_color = "#e2e8f0" if is_modern else "#d1d5db"
        header_align = "left" if is_modern else "center"

        profile = resume_data.get("profile", {})
        contacts = []
        if profile.get("email"):
            contacts.append(profile["email"])
        if profile.get("phone"):
            contacts.append(profile["phone"])
        if profile.get("location"):
            contacts.append(profile["location"])
        if profile.get("linkedin_url"):
            contacts.append(profile["linkedin_url"].replace("https://", ""))
        if profile.get("github_url"):
            contacts.append(profile["github_url"].replace("https://", ""))

        html = f"""
        <div style="font-family: {font_family}; color: #1e293b; background: #ffffff; padding: 32px 36px; line-height: 1.5; font-size: 13.5px; max-width: 820px; margin: 0 auto; box-shadow: 0 4px 6px -1px rgba(0,0,0,0.05); border: 1px solid #f1f5f9; border-radius: 4px;">
            <!-- HEADER -->
            <div style="text-align: {header_align}; margin-bottom: 16px;">
                <h1 style="font-size: 24px; font-weight: 700; color: {primary_color}; margin: 0 0 4px 0; letter-spacing: -0.02em;">{profile.get('full_name', 'Alex Rivera')}</h1>
                {f'<p style="font-size: 13.5px; font-weight: 600; color: {accent_color}; margin: 0 0 6px 0;">{profile.get("headline", "")}</p>' if profile.get("headline") else ''}
                <div style="font-size: 12px; color: #64748b; display: flex; flex-wrap: wrap; justify-content: {header_align if header_align == 'left' else 'center'}; gap: 8px;">
                    {' &bull; '.join(contacts)}
                </div>
            </div>
            <hr style="border: 0; border-top: 1.5px solid {primary_color if is_modern else '#94a3b8'}; margin: 12px 0 16px 0;" />
        """

        # Summary
        summary = resume_data.get("summary") or profile.get("summary")
        if summary:
            html += f"""
            <div style="margin-bottom: 16px;">
                <h2 style="font-size: 13px; font-weight: 700; text-transform: uppercase; letter-spacing: 0.05em; color: {primary_color}; margin: 0 0 6px 0; border-bottom: 1px solid {border_color}; padding-bottom: 3px;">Professional Summary</h2>
                <p style="margin: 0; color: #334155; line-height: 1.6;">{summary}</p>
            </div>
            """

        # Technical Skills
        skills = resume_data.get("skills", [])
        if skills:
            tech_s = [s["name"] for s in skills if s.get("category") == "technical"]
            tool_s = [s["name"] for s in skills if s.get("category") == "tool"]
            soft_s = [s["name"] for s in skills if s.get("category") == "soft"]

            html += f"""
            <div style="margin-bottom: 16px;">
                <h2 style="font-size: 13px; font-weight: 700; text-transform: uppercase; letter-spacing: 0.05em; color: {primary_color}; margin: 0 0 6px 0; border-bottom: 1px solid {border_color}; padding-bottom: 3px;">Technical Skills</h2>
                <div style="display: flex; flex-direction: column; gap: 4px; font-size: 13px;">
                    {f'<div><strong>Languages & Frameworks:</strong> {", ".join(tech_s)}</div>' if tech_s else ''}
                    {f'<div><strong>Tools & Infrastructure:</strong> {", ".join(tool_s)}</div>' if tool_s else ''}
                    {f'<div><strong>Architecture & Leadership:</strong> {", ".join(soft_s)}</div>' if soft_s else ''}
                </div>
            </div>
            """

        # Experience
        experience = resume_data.get("experience", [])
        if experience:
            html += f"""
            <div style="margin-bottom: 16px;">
                <h2 style="font-size: 13px; font-weight: 700; text-transform: uppercase; letter-spacing: 0.05em; color: {primary_color}; margin: 0 0 8px 0; border-bottom: 1px solid {border_color}; padding-bottom: 3px;">Professional Experience</h2>
            """
            for exp in experience:
                bullets_html = "".join([f"<li style='margin-bottom: 4px;'>{b}</li>" for b in exp.get("bullet_points", [])])
                html += f"""
                <div style="margin-bottom: 12px;">
                    <div style="display: flex; justify-content: space-between; align-items: baseline; margin-bottom: 2px;">
                        <div>
                            <strong style="font-size: 13.5px; color: #0f172a;">{exp.get('role')}</strong> &mdash; 
                            <span style="color: #334155; font-weight: 500;">{exp.get('company')}</span>
                        </div>
                        <div style="font-size: 12px; color: #64748b;">{exp.get('start_date')} – {exp.get('end_date', 'Present')}</div>
                    </div>
                    <ul style="margin: 4px 0 0 0; padding-left: 20px; color: #334155; font-size: 12.8px; line-height: 1.55;">
                        {bullets_html}
                    </ul>
                </div>
                """
            html += "</div>"

        # Projects
        projects = resume_data.get("projects", [])
        if projects:
            html += f"""
            <div style="margin-bottom: 16px;">
                <h2 style="font-size: 13px; font-weight: 700; text-transform: uppercase; letter-spacing: 0.05em; color: {primary_color}; margin: 0 0 8px 0; border-bottom: 1px solid {border_color}; padding-bottom: 3px;">Key Projects</h2>
            """
            for p in projects:
                tech_str = f" <span style='font-size: 12px; color: #64748b; font-weight: normal;'>({', '.join(p.get('tech_stack', []))})</span>" if p.get('tech_stack') else ""
                bullets_html = "".join([f"<li style='margin-bottom: 4px;'>{b}</li>" for b in p.get("bullet_points", [])])
                html += f"""
                <div style="margin-bottom: 12px;">
                    <div style="display: flex; justify-content: space-between; align-items: baseline; margin-bottom: 2px;">
                        <div>
                            <strong style="font-size: 13.5px; color: #0f172a;">{p.get('title')}</strong>{tech_str}
                        </div>
                        <div style="font-size: 12px; color: #64748b;">{p.get('start_date', '')} – {p.get('end_date', 'Present')}</div>
                    </div>
                    <ul style="margin: 4px 0 0 0; padding-left: 20px; color: #334155; font-size: 12.8px; line-height: 1.55;">
                        {bullets_html}
                    </ul>
                </div>
                """
            html += "</div>"

        # Education
        education = resume_data.get("education", [])
        if education:
            html += f"""
            <div style="margin-bottom: 16px;">
                <h2 style="font-size: 13px; font-weight: 700; text-transform: uppercase; letter-spacing: 0.05em; color: {primary_color}; margin: 0 0 8px 0; border-bottom: 1px solid {border_color}; padding-bottom: 3px;">Education</h2>
            """
            for edu in education:
                html += f"""
                <div style="display: flex; justify-content: space-between; align-items: baseline; margin-bottom: 4px;">
                    <div>
                        <strong style="color: #0f172a;">{edu.get('institution')}</strong> &mdash; 
                        <span style="color: #334155;">{edu.get('degree')} in {edu.get('field_of_study')}</span>
                        {f' <span style="font-size: 12px; color: #64748b;">(GPA: {edu.get("gpa")})</span>' if edu.get("gpa") else ''}
                    </div>
                    <div style="font-size: 12px; color: #64748b;">{edu.get('start_date')} – {edu.get('end_date', 'Present')}</div>
                </div>
                """
            html += "</div>"

        # Certifications
        certs = resume_data.get("certifications", [])
        if certs:
            html += f"""
            <div style="margin-bottom: 12px;">
                <h2 style="font-size: 13px; font-weight: 700; text-transform: uppercase; letter-spacing: 0.05em; color: {primary_color}; margin: 0 0 6px 0; border-bottom: 1px solid {border_color}; padding-bottom: 3px;">Certifications</h2>
                <ul style="margin: 0; padding-left: 20px; color: #334155; font-size: 12.8px;">
            """
            for c in certs:
                html += f"<li style='margin-bottom: 2px;'><strong>{c.get('name')}</strong> &mdash; {c.get('issuer')} ({c.get('issue_date')})</li>"
            html += "</ul></div>"

        html += "</div>"
        return html
